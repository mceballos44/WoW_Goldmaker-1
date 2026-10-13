"""AH investment, opportunity scanner and Excel export/import (synthetic data only)."""

from datetime import datetime, timezone

import pandas as pd
import pytest
from openpyxl import load_workbook
import io

from goldmaker import conversion, crafting, database, excel_io, investment, pricing, scanner
from goldmaker.currency import gold_to_copper as G
from goldmaker.pricing import PriceResult
from goldmaker.recipes import Reagent, Recipe

NOW = datetime(2026, 10, 13, tzinfo=timezone.utc)


# --- investment --------------------------------------------------------------------------------

def test_investment_baseline_numbers():
    scen = investment.default_scenarios(G(150), sell_through=0.8, spread=0.2)
    r = investment.calculate_investment(G(100), 100, scen, ah_cut=0.05, other_costs_copper=G(50))
    assert r.investment_cost == G(10050)
    base = r.scenario("Baseline")
    assert base.units_sold == 80 and base.gross == G(12000)
    assert base.fees == pytest.approx(G(600)) and base.net_revenue == pytest.approx(G(11400))
    assert base.profit == pytest.approx(G(1350)) and base.roi == pytest.approx(1350 / 10050)
    assert base.breakeven_resale_price == pytest.approx(G(10050) / (80 * 0.95))
    assert r.scenario("Bearish").resale_price_copper == pytest.approx(G(120))
    assert r.scenario("Bullish").resale_price_copper == pytest.approx(G(180))
    assert r.scenario("Bearish").profit < base.profit < r.scenario("Bullish").profit


def test_investment_rejects_missing_or_bad_inputs():
    scen = investment.default_scenarios(G(150), 1.0)
    with pytest.raises(ValueError):
        investment.calculate_investment(None, 10, scen)
    with pytest.raises(ValueError):
        investment.calculate_investment(G(1), 0, scen)
    with pytest.raises(ValueError):
        investment.calculate_investment(G(1), 1, [investment.Scenario("X", G(1), 1.5)])
    with pytest.raises(ValueError):
        investment.default_scenarios(G(1), 1.0, spread=1.0)


def test_zero_sell_through_has_no_breakeven():
    r = investment.calculate_investment(G(1), 1, [investment.Scenario("Dead", G(10), 0.0)])
    assert r.scenarios[0].breakeven_resale_price is None and r.scenarios[0].profit == -G(1)


# --- scanner -----------------------------------------------------------------------------------

def price_of_factory(prices, stale=()):
    def price_of(item_id):
        copper = prices.get(item_id)
        return PriceResult(item_id, f"Item {item_id}", copper, "commodities" if copper is not None else "missing",
                           False, item_id in stale, NOW)
    return price_of


def test_conversion_opportunity_per_operation_and_hour():
    summary = conversion.estimate_yields([conversion.Batch(950, 100, {951: 20, 952: 100})])
    price_of = price_of_factory({950: G(20), 951: G(100), 952: G(10)})
    opp = scanner.from_conversion("prospecting", 950, "Test Ore", summary, price_of, lambda i: (1000.0, None),
                                  operations=10, units_per_operation=5, seconds_per_op=2.0, ah_cut=0.0)
    # per ore: 0.2*100 + 1*10 = 30g, cost 20g -> 10g; per op of 5 ore -> 50g
    assert opp.profit_per_op == pytest.approx(G(50)) and opp.cost_per_op == G(100)
    assert opp.investment == G(1000) and opp.profit_per_hour == pytest.approx(G(50) * 1800)
    assert any("small sample" in f for f in opp.flags)
    assert opp.liquidity_days == pytest.approx(50 / 1000)


def test_conversion_opportunity_missing_price_and_no_size():
    summary = conversion.estimate_yields([conversion.Batch(950, 10, {951: 2})])
    opp = scanner.from_conversion("prospecting", 950, None, summary, price_of_factory({950: G(1)}),
                                  lambda i: (None, None), operations=1)
    assert opp.profit_per_op is None and not opp.complete and opp.profit_per_hour is None
    assert any(f.startswith("no price") for f in opp.flags)
    assert any("operation size not set" in f for f in opp.flags)
    empty = conversion.estimate_yields([])
    assert scanner.from_conversion("milling", 1, None, empty, price_of_factory({}), lambda i: (None, None), 1) is None


def test_recipe_opportunities_and_ranking():
    recipe = Recipe(key="r", name="Thing", crafted_item_id=900, reagents=(Reagent(1, 901, 2),))
    unverified = Recipe(key="u", name="Unpriced", crafted_item_id=910, reagents=(Reagent(1, 911, 1),),
                        last_verified="2026-10-01")
    price_of = price_of_factory({900: G(30), 901: G(10)})
    evals = crafting.evaluate_recipes([recipe, unverified], price_of, lambda i: (2.0, 0.5), ah_cut=0.0)
    opps = scanner.from_recipes(evals, operations=4, seconds_per_craft=None)
    ranked = scanner.rank(opps)
    assert ranked[0].name == "Thing" and ranked[0].profit_per_op == pytest.approx(G(10))
    assert ranked[0].liquidity_days == pytest.approx(2.0) and ranked[0].reliability == "recipe unverified"
    assert ranked[1].profit_per_op is None and ranked[1].reliability == "recipe verified"
    assert scanner.rank(opps, "Profit per hour")[0].profit_per_hour is None     # unknown time stays unknown


# --- Excel -------------------------------------------------------------------------------------

def seeded(conn):
    database.save_recipe(conn, Recipe(key="vial", name="Gold Sunglass Vial", crafted_item_id=900,
                                      reagents=(Reagent(1, 901, 5, "Gold Glass"), Reagent(2, 902, 1, "Gold Stone"))))
    database.upsert_override(conn, 902, G(3.29), "user price", "Gold Stone")
    database.add_milling_experiment(conn, 1, 40, [{"item_id": 2, "name": "Roseate", "quantity": 500},
                                                  {"item_id": 3, "name": "Sallow", "quantity": 100}],
                                    herb_name="Felwort", performed_at="2026-10-13")
    database.add_experiment(conn, "prospecting", 950, 100, [{"item_id": 951, "name": "Gem", "quantity": 20}],
                            input_name="Test Ore")


def price_table(conn):
    prices = pd.DataFrame({"item_id": [900, 901, 2, 3], "name": ["Vial", "Gold Glass", "Roseate", "Sallow"],
                           "market_value": [G(300), G(104), G(2), G(40)], "min_buyout": None, "recent": None,
                           "historical": None, "avg_sale_price": None, "sale_rate": None, "sold_per_day": None,
                           "scope": "commodities", "dataset_key": "us-commodities",
                           "source_updated_at": NOW.isoformat(), "fetched_at": NOW.isoformat()})
    return pricing.build_price_table(prices, database.list_overrides(conn), now=NOW)


def test_full_export_has_sheets_formulas_and_blank_missing(conn):
    seeded(conn)
    table = price_table(conn)
    data = excel_io.collect(conn, table, {"Region": "us", "AH cut %": 5}, ah_cut=0.05)
    data.investment = investment.calculate_investment(G(100), 10, investment.default_scenarios(G(150), 1.0))
    wb = load_workbook(io.BytesIO(excel_io.build_workbook(data)))
    assert {"Info", "Price snapshot", "Prices", "Recipes", "Crafting profitability", "Material breakdown",
            "Experiments", "Yields", "Investment"} <= set(wb.sheetnames)
    craft = wb["Crafting profitability"]
    assert craft["A2"].value == "Gold Sunglass Vial"
    assert craft["B2"].value == pytest.approx(523.29) and craft["C2"].value == pytest.approx(300)
    assert craft["H2"].value.startswith("=IF(")                   # profit is a live formula
    yields = {(r[0], r[3]): r for r in wb["Yields"].iter_rows(min_row=2, values_only=True)}
    assert yields[("milling", "Roseate")][5] == pytest.approx(12.5)
    assert yields[("prospecting", "Gem")][12] is None              # no price -> blank, not 0
    assert wb["Investment"]["B7"].value == "=B3*B4+B5"


def test_export_import_round_trip_into_fresh_database(conn, tmp_path):
    seeded(conn)
    content = excel_io.build_workbook(excel_io.collect(conn, price_table(conn), {}))
    parsed = excel_io.read_workbook(content)
    assert parsed.errors == [] and parsed.recipes.errors == []
    assert [r.key for r in parsed.recipes.recipes] == ["vial"]
    assert parsed.prices == [{"item_id": 902, "name": "Gold Stone", "price_copper": G(3.29), "note": "user price"}]
    assert {e["process"] for e in parsed.experiments} == {"milling", "prospecting"}

    fresh = database.connect(tmp_path / "fresh.sqlite3")
    try:
        counts = excel_io.apply_import(fresh, parsed)
        assert counts == {"recipes": 1, "prices": 1, "experiments": 2, "experiments_skipped": 0}
        again = excel_io.apply_import(fresh, parsed)               # same file twice: no double counting
        assert again["experiments"] == 0 and again["experiments_skipped"] == 2
        felwort = database.list_milling_experiments(fresh)
        assert {o["item_id"]: o["quantity"] for o in felwort[0]["outputs"]} == {2: 500, 3: 100}
    finally:
        fresh.close()


def test_template_and_bad_rows():
    from openpyxl import Workbook
    template = load_workbook(io.BytesIO(excel_io.template_workbook()))
    assert [c.value for c in template["Experiments"][1]] == excel_io.EXPERIMENT_COLUMNS
    assert excel_io.read_workbook(excel_io.template_workbook()).empty

    wb = Workbook()
    ws = wb.active
    ws.title = "Prices"
    ws.append(excel_io.PRICE_SHEET_COLUMNS)
    ws.append([5, "A", None, ""])          # missing price -> error, never 0
    ws.append([6, "B", -1, ""])
    ws.append([7, "C", 2.5, "ok"])
    exp = wb.create_sheet("Experiments")
    exp.append(excel_io.EXPERIMENT_COLUMNS)
    exp.append(["x", "smelting", 1, None, 5, None, None, 2, None, 1])
    buffer = io.BytesIO()
    wb.save(buffer)
    parsed = excel_io.read_workbook(buffer.getvalue())
    assert parsed.prices == [{"item_id": 7, "name": "C", "price_copper": G(2.5), "note": "ok"}]
    text = " ".join(parsed.errors)
    assert "row 2" in text and "row 3" in text and "process must be one of" in text
    assert excel_io.read_workbook(b"not a workbook").errors
