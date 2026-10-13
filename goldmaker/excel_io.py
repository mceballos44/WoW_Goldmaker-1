"""Excel export and import (openpyxl).

Export: one workbook with formatted sheets - an *Info* sheet (time, region/realm, price
basis, AH cut, TSM data timestamps, assumptions), price snapshot, manual prices,
recipes and their material breakdown, crafting profitability, experiments, calibrated
yields, opportunities and investment scenarios. Profit columns are Excel FORMULAS, so
changing a price in Excel recalculates the row. A missing price is left blank and the
formulas stay blank too - it never becomes 0.

Import: the sheets *Recipes*, *Prices* (manual prices) and *Experiments* use the same
columns the export writes, so an exported workbook can be edited and read back.
``template_workbook`` gives empty sheets with those headers.

Money is COPPER in this app; sheets show GOLD.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from typing import Optional

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from . import config, conversion, crafting, database, pricing, recipes
from .currency import copper_to_gold, gold_to_copper, is_missing

PRICE_SHEET_COLUMNS = ["item_id", "name", "price_gold", "note"]
EXPERIMENT_COLUMNS = ["experiment_ref", "process", "input_item_id", "input_name", "units_used",
                      "performed_at", "notes", "output_item_id", "output_name", "quantity"]
GOLD_FORMAT = "#,##0.00"
PCT_FORMAT = "0.0%"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
ASSUMPTIONS = (
    "Prices are TSM periodic snapshots (see the data times below), not real-time.",
    "Manual price overrides win over TSM prices; they are listed on the 'Prices' sheet.",
    "AH commission is one percentage of each sale; per-auction rounding and deposits are ignored.",
    "Revenue is linear: selling more does not move the price. Unsold units are worth 0.",
    "Yields come only from recorded experiments; reliability labels are rule-of-thumb.",
    "Recipes and quantities come only from the user's imports/edits.",
    "Profits are theoretical margins at listed prices, not demonstrated sales.",
)


@dataclass
class ExportData:
    context: dict = field(default_factory=dict)          # region, realm, basis, AH cut % ...
    datasets: Optional[pd.DataFrame] = None
    price_table: Optional[pd.DataFrame] = None           # pricing.build_price_table output
    overrides: Optional[pd.DataFrame] = None
    recipes: list = field(default_factory=list)
    craft_rows: list = field(default_factory=list)
    experiments: list = field(default_factory=list)      # database.list_experiments(conn, None)
    yields: list = field(default_factory=list)
    opportunities: list = field(default_factory=list)
    investment: Optional[object] = None                  # investment.InvestmentResult
    investment_name: str = ""
    generated_at: datetime = field(default_factory=lambda: datetime.now(timezone.utc))


# --- collecting data for a full export ------------------------------------------------------------

def collect(conn, table: pd.DataFrame, context: dict, ah_cut: float = 0.05,
            sell_through: float = 1.0, opportunities: Optional[list] = None) -> ExportData:
    """Gather everything stored locally into an ``ExportData`` (prices from ``table``)."""
    index = pricing.PriceIndex(table)
    recipe_list = [recipes.Recipe.from_row(r) for r in database.list_recipe_rows(conn)]
    resources = database.list_item_resources(conn)
    opportunity = {int(r.item_id): int(r.opportunity_cost_copper)
                   for r in resources.itertuples() if pd.notna(r.opportunity_cost_copper)}
    on_hand = {int(r.item_id): int(r.on_hand) for r in resources.itertuples() if pd.notna(r.on_hand)}
    evaluations = crafting.evaluate_recipes(recipe_list, index.get, index.demand, opportunity, on_hand,
                                            ah_cut=ah_cut, sell_through=sell_through)
    craft_rows = [{
        "recipe": e.recipe.name, "key": e.recipe.key, "profession": e.recipe.profession,
        "patch": e.recipe.game_patch, "last_verified": e.recipe.last_verified,
        "source": e.recipe.source, "cost": e.result.cost_per_craft,
        "output_price": e.result.output_price_copper, "output_units": e.result.output_units,
        "ah_cut": ah_cut, "sell_through": sell_through, "flags": "; ".join(e.flags),
    } for e in evaluations]

    experiments = database.list_experiments(conn, None)
    yield_rows = []
    for process in database.PROCESSES:
        inputs = database.list_experiment_inputs(conn, process)
        for row in inputs.itertuples():
            batches = [conversion.Batch.from_row(e) for e in experiments
                       if e["process"] == process and e["herb_item_id"] == row.herb_item_id]
            summary = conversion.estimate_yields(batches)
            for est in summary.estimates:
                price = index.get(est.item_id)
                yield_rows.append({
                    "process": process, "input_item_id": int(row.herb_item_id),
                    "input_name": row.herb_name, "output_item_id": est.item_id, "output_name": est.name,
                    "mean_per_unit": est.mean_per_unit, "std_error": est.std_error,
                    "ci_low": est.ci_low, "ci_high": est.ci_high,
                    "experiments": summary.n_experiments, "total_units": summary.total_units,
                    "reliability": summary.reliability, "price": price.copper})
    return ExportData(
        context=context, datasets=database.list_datasets(conn), price_table=table,
        overrides=database.list_overrides(conn), recipes=recipe_list, craft_rows=craft_rows,
        experiments=experiments, yields=yield_rows, opportunities=opportunities or [])


# --- writing --------------------------------------------------------------------------------------

def _g(copper):
    return copper_to_gold(copper)


def _cell(value):
    """Excel-safe value: NaN/None -> blank, timestamps -> naive UTC datetimes."""
    if value is None or (not isinstance(value, str) and is_missing(value)):
        return None
    if isinstance(value, pd.Timestamp):
        value = value.to_pydatetime()
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.astimezone(timezone.utc).replace(tzinfo=None)
    if hasattr(value, "item"):
        return value.item()
    return value


def _table(ws, name: str, headers: list, rows: list, formats: Optional[dict] = None,
           start_row: int = 1) -> None:
    """Write a header + rows as a formatted Excel table."""
    formats = formats or {}
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=header)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r, row in enumerate(rows, start=start_row + 1):
        for col, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=col, value=_cell(value))
            if headers[col - 1] in formats:
                cell.number_format = formats[headers[col - 1]]
    for col, header in enumerate(headers, start=1):
        width = max([len(str(header))] + [len(str(_cell(row[col - 1]) or "")) for row in rows[:200]])
        ws.column_dimensions[get_column_letter(col)].width = min(max(width + 2, 10), 60)
    if rows:
        ref = f"A{start_row}:{get_column_letter(len(headers))}{start_row + len(rows)}"
        table = Table(displayName=name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)


def _info_sheet(wb: Workbook, data: ExportData) -> None:
    ws = wb.active
    ws.title = "Info"
    ws["A1"] = "WoW Goldmaker export"
    ws["A1"].font = Font(bold=True, size=14)
    rows = [("Generated (UTC)", _cell(data.generated_at))]
    rows += [(str(k), _cell(v)) for k, v in data.context.items()]
    for r, (label, value) in enumerate(rows, start=3):
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        ws.cell(row=r, column=2, value=value)
    r = len(rows) + 4
    ws.cell(row=r, column=1, value="Assumptions").font = Font(bold=True)
    for offset, text in enumerate(ASSUMPTIONS, start=1):
        ws.cell(row=r + offset, column=1, value=f"- {text}")
    r += len(ASSUMPTIONS) + 2
    if data.datasets is not None and not data.datasets.empty:
        ws.cell(row=r, column=1, value="TSM datasets used").font = Font(bold=True)
        cols = ["dataset_key", "scope", "source_updated_at", "fetched_at", "row_count", "origin"]
        frame = data.datasets.reindex(columns=cols)
        _table(ws, "Datasets", cols, frame.values.tolist(), start_row=r + 1)
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 30
    ws.freeze_panes = None


def _price_sheets(wb: Workbook, data: ExportData) -> None:
    table = data.price_table
    if table is not None and not table.empty:
        ws = wb.create_sheet("Price snapshot")
        headers = ["Item ID", "Name", "Effective price (g)", "Price source", "TSM price (g)",
                   "Override (g)", "Sold / day", "Sale rate", "TSM data time (UTC)", "Stale?"]
        rows = [[r.item_id, r.name, _g(r.effective_copper), r.effective_source, _g(r.tsm_copper),
                 _g(r.override_copper), r.sold_per_day, r.sale_rate, r.updated_at, bool(r.tsm_stale)]
                for r in table.itertuples()]
        _table(ws, "PriceSnapshot", headers, rows, {
            "Effective price (g)": GOLD_FORMAT, "TSM price (g)": GOLD_FORMAT, "Override (g)": GOLD_FORMAT,
            "TSM data time (UTC)": "yyyy-mm-dd hh:mm"})
    ws = wb.create_sheet("Prices")
    overrides = data.overrides if data.overrides is not None else pd.DataFrame()
    rows = [[r.item_id, r.name, _g(r.price_copper), r.note] for r in overrides.itertuples()]
    _table(ws, "ManualPrices", PRICE_SHEET_COLUMNS, rows, {"price_gold": GOLD_FORMAT})


def _recipe_sheets(wb: Workbook, data: ExportData) -> None:
    ws = wb.create_sheet("Recipes")
    rows = [[row[c] for c in recipes.CSV_COLUMNS] for row in recipes.flat_rows(data.recipes)]
    _table(ws, "Recipes", recipes.CSV_COLUMNS, rows)
    if not data.craft_rows:
        return
    ws = wb.create_sheet("Crafting profitability")
    headers = ["Recipe", "Cost / craft (g)", "Output price (g)", "Output units / craft", "AH cut",
               "Sell-through", "Net / craft (g)", "Profit / craft (g)", "ROI", "Break-even price (g)",
               "Profession", "Patch", "Last verified", "Source", "Flags", "Recipe key"]
    rows = []
    for n, c in enumerate(data.craft_rows, start=2):
        rows.append([
            c["recipe"], _g(c["cost"]), _g(c["output_price"]), c["output_units"], c["ah_cut"],
            c["sell_through"],
            f'=IF(OR(C{n}="",D{n}=""),"",C{n}*D{n}*(1-E{n})*F{n})',
            f'=IF(OR(B{n}="",G{n}=""),"",G{n}-B{n})',
            f'=IF(OR(H{n}="",B{n}="",B{n}=0),"",H{n}/B{n})',
            f'=IF(OR(B{n}="",D{n}*(1-E{n})*F{n}=0),"",B{n}/(D{n}*(1-E{n})*F{n}))',
            c["profession"], c["patch"], c["last_verified"] or "never", c["source"], c["flags"], c["key"]])
    _table(ws, "Crafting", headers, rows, {
        "Cost / craft (g)": GOLD_FORMAT, "Output price (g)": GOLD_FORMAT, "AH cut": PCT_FORMAT,
        "Sell-through": PCT_FORMAT, "Net / craft (g)": GOLD_FORMAT, "Profit / craft (g)": GOLD_FORMAT,
        "ROI": PCT_FORMAT, "Break-even price (g)": GOLD_FORMAT})

    ws = wb.create_sheet("Material breakdown")
    headers = ["Recipe", "Slot", "Reagent", "Item ID", "Quantity", "Quality", "Auctionable",
               "Savings eligible", "Unit price (g)", "Line cost (g)"]
    prices = {}
    if data.price_table is not None and not data.price_table.empty:
        prices = dict(zip(data.price_table["item_id"], data.price_table["effective_copper"]))
    rows = []
    for recipe in data.recipes:
        for g in recipe.reagents:
            n = len(rows) + 2
            price = prices.get(g.item_id)
            rows.append([recipe.name, g.slot, g.item_name, g.item_id, g.quantity, g.quality_label,
                         g.auctionable, g.savings_eligible, _g(price) if g.auctionable else None,
                         f'=IF(I{n}="","",E{n}*I{n})'])
    _table(ws, "Materials", headers, rows, {"Unit price (g)": GOLD_FORMAT, "Line cost (g)": GOLD_FORMAT})


def _experiment_sheets(wb: Workbook, data: ExportData) -> None:
    ws = wb.create_sheet("Experiments")
    rows = []
    for e in data.experiments:
        for o in e["outputs"] or [{"item_id": None, "item_name": None, "quantity": None}]:
            rows.append([e["id"], e.get("process", "milling"), e["herb_item_id"], e["herb_name"],
                         e["herbs_used"], e["performed_at"], e["notes"], o["item_id"],
                         o["item_name"], o["quantity"]])
    _table(ws, "Experiments", EXPERIMENT_COLUMNS, rows)
    if not data.yields:
        return
    ws = wb.create_sheet("Yields")
    headers = ["Process", "Input", "Input ID", "Output", "Output ID", "Mean per input unit", "Std. error",
               "95% low", "95% high", "Experiments", "Input units", "Reliability", "Output price (g)",
               "Gross value per input unit (g)"]
    rows = []
    for n, y in enumerate(data.yields, start=2):
        rows.append([y["process"], y["input_name"], y["input_item_id"], y["output_name"],
                     y["output_item_id"], y["mean_per_unit"], y["std_error"], y["ci_low"], y["ci_high"],
                     y["experiments"], y["total_units"], y["reliability"], _g(y["price"]),
                     f'=IF(M{n}="","",F{n}*M{n})'])
    _table(ws, "Yields", headers, rows, {"Output price (g)": GOLD_FORMAT,
                                          "Gross value per input unit (g)": GOLD_FORMAT})


def _opportunity_sheet(wb: Workbook, data: ExportData) -> None:
    if not data.opportunities:
        return
    ws = wb.create_sheet("Opportunities")
    headers = ["Kind", "Name", "Operation", "Cost / op (g)", "Profit / op (g)", "ROI", "Profit / hour (g)",
               "Investment (g)", "Days of daily volume", "Reliability", "Flags"]
    rows = [[o.kind, o.name, o.operation, _g(o.cost_per_op), _g(o.profit_per_op), o.roi,
             _g(o.profit_per_hour), _g(o.investment), o.liquidity_days, o.reliability, "; ".join(o.flags)]
            for o in data.opportunities]
    _table(ws, "Opportunities", headers, rows, {
        "Cost / op (g)": GOLD_FORMAT, "Profit / op (g)": GOLD_FORMAT, "ROI": PCT_FORMAT,
        "Profit / hour (g)": GOLD_FORMAT, "Investment (g)": GOLD_FORMAT, "Days of daily volume": "0.00"})


def _investment_sheet(wb: Workbook, data: ExportData) -> None:
    inv = data.investment
    if inv is None:
        return
    ws = wb.create_sheet("Investment")
    ws["A1"] = f"AH investment{': ' + data.investment_name if data.investment_name else ''}"
    ws["A1"].font = Font(bold=True, size=13)
    inputs = [("Buy price (g each)", _g(inv.buy_price_copper), GOLD_FORMAT),
              ("Quantity", inv.quantity, "0"),
              ("Other costs: deposits lost, re-listing (g)", _g(inv.other_costs_copper), GOLD_FORMAT),
              ("AH cut", inv.ah_cut, PCT_FORMAT)]
    for r, (label, value, fmt) in enumerate(inputs, start=3):
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        ws.cell(row=r, column=2, value=value).number_format = fmt
    ws.cell(row=7, column=1, value="Investment cost (g)").font = Font(bold=True)
    ws.cell(row=7, column=2, value="=B3*B4+B5").number_format = GOLD_FORMAT
    headers = ["Scenario", "Resale price (g)", "Sell-through", "Units sold", "Gross (g)", "AH fees (g)",
               "Net revenue (g)", "Profit (g)", "ROI", "Break-even resale price (g)"]
    rows = []
    for n, s in enumerate(inv.scenarios, start=10):
        rows.append([s.name, _g(s.resale_price_copper), s.sell_through, f"=$B$4*C{n}", f"=D{n}*B{n}",
                     f"=E{n}*$B$6", f"=E{n}-F{n}", f"=G{n}-$B$7", f'=IF($B$7=0,"",H{n}/$B$7)',
                     f'=IF(D{n}*(1-$B$6)=0,"",$B$7/(D{n}*(1-$B$6)))'])
    _table(ws, "Scenarios", headers, rows, {
        "Resale price (g)": GOLD_FORMAT, "Sell-through": PCT_FORMAT, "Units sold": "0.0",
        "Gross (g)": GOLD_FORMAT, "AH fees (g)": GOLD_FORMAT, "Net revenue (g)": GOLD_FORMAT,
        "Profit (g)": GOLD_FORMAT, "ROI": PCT_FORMAT, "Break-even resale price (g)": GOLD_FORMAT},
        start_row=9)
    ws.column_dimensions["A"].width = 42
    ws.cell(row=14, column=1, value="Edit the inputs in column B or a scenario's price / sell-through; "
                                    "the formulas recalculate. Future prices are your own estimates.")


def build_workbook(data: ExportData) -> bytes:
    """The export as .xlsx bytes. Sheets without data are left out (Info and Prices always exist)."""
    wb = Workbook()
    _info_sheet(wb, data)
    _price_sheets(wb, data)
    if data.recipes:
        _recipe_sheets(wb, data)
    if data.experiments:
        _experiment_sheets(wb, data)
    _opportunity_sheet(wb, data)
    _investment_sheet(wb, data)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def template_workbook() -> bytes:
    """Empty Recipes / Prices / Experiments sheets with the import headers."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Info"
    lines = ["WoW Goldmaker import template",
             "Fill in any of the sheets 'Recipes', 'Prices', 'Experiments' and import the file.",
             "Recipes: one row per reagent - same columns as the recipe CSV format (see README).",
             "Prices: manual price overrides in GOLD per unit; they win over TSM prices.",
             "Experiments: one row per output item; rows sharing an experiment_ref are one batch.",
             "process is 'milling' or 'prospecting'. units_used = herbs or ore used in that batch.",
             "Use your own measured numbers only. Leave a price blank if you do not know it."]
    for r, text in enumerate(lines, start=1):
        ws.cell(row=r, column=1, value=text).font = Font(bold=(r == 1))
    for title, headers in (("Recipes", recipes.CSV_COLUMNS), ("Prices", PRICE_SHEET_COLUMNS),
                           ("Experiments", EXPERIMENT_COLUMNS)):
        sheet = wb.create_sheet(title)
        _table(sheet, title, headers, [])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# --- reading --------------------------------------------------------------------------------------

@dataclass
class WorkbookImport:
    recipes: Optional[recipes.ImportResult] = None
    prices: list = field(default_factory=list)          # {"item_id", "name", "price_copper", "note"}
    experiments: list = field(default_factory=list)     # kwargs for database.add_experiment
    errors: list = field(default_factory=list)
    sheets_found: list = field(default_factory=list)

    @property
    def empty(self) -> bool:
        return not (self.prices or self.experiments or (self.recipes and self.recipes.recipes))


def _plain(value):
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def _sheet_rows(ws) -> list:
    """[(excel_row_number, {lower-case header: value})] for non-empty rows."""
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
    out = []
    for number, values in enumerate(rows[1:], start=2):
        if all(v is None or str(v).strip() == "" for v in values):
            continue
        out.append((number, {h: _plain(v) for h, v in zip(headers, values) if h}))
    return out


def _number(value, label: str, errors: list, whole: bool = False):
    if value is None or str(value).strip() == "":
        errors.append(f"{label} is missing")
        return None
    try:
        number = float(str(value).replace(",", ""))
    except ValueError:
        errors.append(f"{label} '{value}' is not a number")
        return None
    if whole:
        if number != int(number):
            errors.append(f"{label} '{value}' must be a whole number")
            return None
        return int(number)
    return number


def read_workbook(content: bytes, lookup: Optional[recipes.Lookup] = None) -> WorkbookImport:
    """Parse the importable sheets. Nothing is written to the database here."""
    result = WorkbookImport()
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # openpyxl raises several types for non-xlsx input
        result.errors.append(f"Could not open the workbook: {exc}")
        return result
    sheets = {name.strip().lower(): wb[name] for name in wb.sheetnames}

    if "recipes" in sheets:
        result.sheets_found.append("Recipes")
        rows = _sheet_rows(sheets["recipes"])
        if rows:
            result.recipes = recipes.parse_recipe_rows([(f"Recipes!{n}", r) for n, r in rows], lookup)

    if "prices" in sheets:
        result.sheets_found.append("Prices")
        for n, row in _sheet_rows(sheets["prices"]):
            errors: list = []
            item_id = _number(row.get("item_id"), "item_id", errors, whole=True)
            gold = _number(row.get("price_gold"), "price_gold", errors)
            if gold is not None and gold < 0:
                errors.append("price_gold cannot be negative")
            if item_id is not None and item_id <= 0:
                errors.append("item_id must be positive")
            if errors:
                result.errors.append(f"Prices row {n}: " + "; ".join(errors))
                continue
            result.prices.append({"item_id": item_id, "name": row.get("name") or None,
                                  "price_copper": gold_to_copper(gold), "note": row.get("note") or ""})

    if "experiments" in sheets:
        result.sheets_found.append("Experiments")
        groups: dict = {}
        for n, row in _sheet_rows(sheets["experiments"]):
            ref = row.get("experiment_ref")
            if ref is None or str(ref).strip() == "":
                result.errors.append(f"Experiments row {n}: experiment_ref is required")
                continue
            groups.setdefault(str(ref).strip(), []).append((n, row))
        for ref, group in groups.items():
            errors = []
            first = group[0][1]
            process = str(first.get("process") or "").strip().lower()
            if process not in database.PROCESSES:
                errors.append(f"process must be one of {', '.join(database.PROCESSES)}")
            input_id = _number(first.get("input_item_id"), "input_item_id", errors, whole=True)
            units = _number(first.get("units_used"), "units_used", errors, whole=True)
            if units is not None and units <= 0:
                errors.append("units_used must be greater than zero")
            outputs = []
            for n, row in group:
                out_id = _number(row.get("output_item_id"), f"row {n} output_item_id", errors, whole=True)
                qty = _number(row.get("quantity"), f"row {n} quantity", errors, whole=True)
                if qty is not None and qty < 0:
                    errors.append(f"row {n} quantity cannot be negative")
                if out_id is not None and qty is not None:
                    outputs.append({"item_id": out_id, "name": row.get("output_name") or None,
                                    "quantity": qty})
            if errors:
                result.errors.append(f"Experiment '{ref}': " + "; ".join(errors))
                continue
            result.experiments.append({
                "process": process, "input_item_id": input_id, "units_used": units, "outputs": outputs,
                "input_name": first.get("input_name") or None,
                "performed_at": str(first.get("performed_at")) if first.get("performed_at") else None,
                "notes": first.get("notes") or f"Imported from Excel (ref {ref})"})
    if not result.sheets_found:
        result.errors.append("No sheet called 'Recipes', 'Prices' or 'Experiments' was found.")
    wb.close()
    return result


def apply_import(conn, data: WorkbookImport, skip_existing_experiments: bool = True) -> dict:
    """Write a parsed workbook to the database. Returns counts of what was saved.

    Experiments identical to one already stored (same process, input, units and outputs)
    are skipped by default so importing the same file twice does not double-count yields.
    """
    counts = {"recipes": 0, "prices": 0, "experiments": 0, "experiments_skipped": 0}
    if data.recipes:
        for recipe in data.recipes.recipes:
            database.save_recipe(conn, recipe)
            counts["recipes"] += 1
    for price in data.prices:
        database.upsert_override(conn, price["item_id"], price["price_copper"], price["note"], price["name"])
        counts["prices"] += 1
    existing = database.list_experiments(conn, None)
    signatures = {(e["process"], e["herb_item_id"], e["herbs_used"],
                   tuple(sorted((o["item_id"], o["quantity"]) for o in e["outputs"]))) for e in existing}
    for exp in data.experiments:
        merged: dict = {}
        for o in exp["outputs"]:
            merged[o["item_id"]] = merged.get(o["item_id"], 0) + o["quantity"]
        signature = (exp["process"], exp["input_item_id"], exp["units_used"], tuple(sorted(merged.items())))
        if skip_existing_experiments and signature in signatures:
            counts["experiments_skipped"] += 1
            continue
        database.add_experiment(conn, **exp)
        signatures.add(signature)
        counts["experiments"] += 1
    return counts
